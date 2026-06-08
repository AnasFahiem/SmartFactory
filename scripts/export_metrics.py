import os
import sys

def install_and_import(package):
    import importlib
    try:
        importlib.import_module(package)
    except ImportError:
        import subprocess
        print(f"Installing {package} for exporting metrics...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])

# Ensure pandas and openpyxl are installed
install_and_import("pandas")
install_and_import("openpyxl")

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def export_training_results(results_csv_path, output_excel_path):
    if not os.path.exists(results_csv_path):
        print(f"Error: Could not find results CSV at {results_csv_path}")
        return False

    print(f"Reading training results from {results_csv_path}...")
    
    # Read YOLO results.csv
    # YOLO results.csv headers often contain leading/trailing whitespaces, so we strip them
    df = pd.read_csv(results_csv_path)
    df.columns = df.columns.str.strip()

    # Create openpyxl workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Training Epoch Metrics"
    
    # Enable grid lines explicitly
    ws.views.sheetView[0].showGridLines = True

    # Style definitions
    font_family = "Segoe UI"
    
    title_font = Font(name=font_family, size=16, bold=True, color="1F4E78")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10)
    best_row_font = Font(name=font_family, size=10, bold=True, color="375623")
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    best_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft green
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Write Title
    ws.append([])
    ws.cell(row=2, column=2, value="SmartFactory PPE Detection Training Metrics").font = title_font
    ws.cell(row=3, column=2, value="Dataset: SH17 (8,099 images, 17 classes) | Model: YOLOv11l").font = Font(name=font_family, size=10, italic=True, color="595959")
    ws.append([])
    ws.append([])

    # Find best epoch based on val mAP50
    # Common column names in YOLO: 'metrics/mAP50(B)' or 'metrics/mAP50(M)'
    map50_col = None
    for col in df.columns:
        if 'mAP50(B)' in col or 'mAP50' in col:
            map50_col = col
            break
            
    best_epoch_idx = -1
    if map50_col and not df[map50_col].empty:
        best_epoch_idx = df[map50_col].idxmax()
        best_map50 = df.loc[best_epoch_idx, map50_col]
        best_epoch_num = df.loc[best_epoch_idx, 'epoch'] if 'epoch' in df.columns else (best_epoch_idx + 1)
        ws.cell(row=4, column=2, value=f"⭐ Best Epoch: Epoch {int(best_epoch_num)} (mAP50: {best_map50:.4f})").font = Font(name=font_family, size=11, bold=True, color="375623")

    ws.append([]) # spacer

    # 2. Write Headers
    start_row = 7
    headers = list(df.columns)
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    # 3. Write Data Rows
    current_row = start_row + 1
    for r_idx, row in df.iterrows():
        is_best_row = (r_idx == best_epoch_idx)
        is_zebra = (r_idx % 2 == 1)
        
        for c_idx, val in enumerate(row, start=2):
            cell = ws.cell(row=current_row, column=c_idx)
            
            # Format numbers beautifully
            if isinstance(val, float):
                cell.value = round(val, 5)
                cell.number_format = '0.00000'
            else:
                cell.value = val
                
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Formatting styles
            if is_best_row:
                cell.fill = best_fill
                cell.font = best_row_font
            elif is_zebra:
                cell.fill = zebra_fill
                cell.font = data_font
            else:
                cell.font = data_font
                
        current_row += 1

    # Freeze panes so headers stay visible when scrolling
    ws.freeze_panes = f"B{start_row + 1}"

    # Auto-adjust column widths
    for col in ws.columns:
        # Ignore columns before col 2 (A is empty)
        if col[0].column < 2:
            continue
        max_len = 0
        for cell in col:
            if cell.row >= start_row and cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save Workbook
    wb.save(output_excel_path)
    print(f"Successfully exported beautifully formatted metrics to: {output_excel_path}")
    return True

if __name__ == "__main__":
    # Test execution
    script_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.dirname(script_dir)
    results_csv = os.path.join(root_dir, "runs", "detect", "sh17_train", "results.csv")
    output_excel = os.path.join(root_dir, "SmartFactory_SH17_Metrics.xlsx")
    export_training_results(results_csv, output_excel)
